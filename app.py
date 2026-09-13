# -*- coding: utf-8 -*-
"""
app.py — "链接转文字" 工作台主服务

用法:
    python app.py                 # 监听 0.0.0.0:8300（手机同 WiFi 可访问）
    PORT=9000 python app.py       # 自定义端口

环境变量:
    WHISPER_MODEL   whisper 模型档位: tiny/base/small/medium/large-v3（默认 medium）
    PORT            监听端口（默认 8300）
    KEEP_HISTORY    是否保留历史记录文件（默认 1）
"""
import io
import json
import logging
import os
import shutil
import tempfile
import threading
import time
import uuid

from fastapi import FastAPI, UploadFile, File
from fastapi.responses import JSONResponse, FileResponse, Response
from pydantic import BaseModel
from fastapi.staticfiles import StaticFiles

import extractor
from cookies_manager import get_cookie_file_for
import transcriber

logging.basicConfig(level=logging.INFO,
                    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s")
log = logging.getLogger("app")

app = FastAPI(title="链接转文字工作台")

# ---------------- 任务存储（内存 + 可选落盘） ----------------
JOBS = {}            # job_id -> dict
JOBS_LOCK = threading.Lock()
WORK_ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "jobs")
HISTORY_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "history")
os.makedirs(WORK_ROOT, exist_ok=True)


def _save_history(job_id: str, result: dict):
    if os.environ.get("KEEP_HISTORY", "1") != "1":
        return
    try:
        os.makedirs(HISTORY_DIR, exist_ok=True)
        rec = {"job_id": job_id, "time": time.strftime("%Y-%m-%d %H:%M:%S"),
               "url": result.get("url", ""), "title": result.get("title", ""),
               "uploader": result.get("uploader", ""),
               "duration": result.get("duration", 0),
               "full_text": result.get("full_text", ""),
               "segments": result.get("segments", [])}
        with open(os.path.join(HISTORY_DIR, f"{job_id}.json"), "w",
                  encoding="utf-8") as f:
            json.dump(rec, f, ensure_ascii=False, indent=1)
    except Exception as e:
        log.warning("保存历史失败: %s", e)


class ExtractRequest(BaseModel):
    text: str            # 用户粘贴的完整分享文本（或纯 URL）


def _run_job(job_id: str, raw_text: str):
    job = JOBS[job_id]
    workdir = os.path.join(WORK_ROOT, job_id)
    os.makedirs(workdir, exist_ok=True)
    try:
        # 1. 提取 URL
        url = extractor.extract_url(raw_text)
        if not url:
            raise ValueError("粘贴的内容里没有找到链接，请复制包含 https:// 的分享文本")
        job.update(stage="解析链接", progress=3, url=url)

        # 2. cookies（抖音/B站自动处理）
        cookie_file = get_cookie_file_for(url)

        # 3. 探测元数据
        info = extractor.probe_info(url, cookie_file)
        job.update(title=info["title"], uploader=info["uploader"],
                   duration=info["duration"])

        # 4. 下载 + 抽音轨
        def dl_progress(p, stage):
            job.update(progress=p, stage=stage)
        wav_path, meta = extractor.download_audio(url, workdir, cookie_file,
                                                  progress_cb=dl_progress)
        job.update(title=meta["title"] or info["title"],
                   uploader=meta["uploader"] or info["uploader"],
                   duration=meta["duration"] or info["duration"])

        # 5. 转写
        def tr_progress(p, stage):
            job.update(progress=p, stage=stage)
        result = transcriber.transcribe(wav_path, progress_cb=tr_progress)

        # 6. 完成
        job.update(
            stage="完成", progress=100, status="done",
            segments=result["segments"], full_text=result["full_text"],
        )
        _save_history(job_id, {**job, "url": url})
    except Exception as e:
        log.exception("任务 %s 失败", job_id)
        job.update(status="error", stage="失败", error=str(e)[:500])
    finally:
        shutil.rmtree(workdir, ignore_errors=True)


@app.post("/api/extract")
def create_job(req: ExtractRequest):
    job_id = uuid.uuid4().hex[:12]
    with JOBS_LOCK:
        JOBS[job_id] = {"id": job_id, "status": "running", "progress": 0,
                        "stage": "排队中", "created": time.time(),
                        "text": req.text[:2000]}
    threading.Thread(target=_run_job, args=(job_id, req.text),
                     daemon=True).start()
    return {"job_id": job_id}


@app.get("/api/status/{job_id}")
def job_status(job_id: str):
    job = JOBS.get(job_id)
    if not job:
        # 尝试从历史恢复（服务重启后仍可查看结果）
        p = os.path.join(HISTORY_DIR, f"{job_id}.json")
        if os.path.exists(p):
            with open(p, encoding="utf-8") as f:
                rec = json.load(f)
            rec.update(status="done", progress=100, stage="完成")
            return rec
        return JSONResponse({"error": "任务不存在"}, status_code=404)
    return {k: v for k, v in job.items() if k != "text"}


# ---------------- SRT 字幕导出 ----------------
def _fmt_srt_ts(t: float) -> str:
    """秒 → SRT 时间戳 00:00:00,000"""
    ms = int(round(t * 1000))
    h, rem = divmod(ms, 3600_000)
    m, rem = divmod(rem, 60_000)
    s, ms = divmod(rem, 1000)
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"


def _make_srt(segments: list) -> str:
    blocks = []
    for i, s in enumerate(segments, 1):
        blocks.append(
            f"{i}\n{_fmt_srt_ts(s['start'])} --> {_fmt_srt_ts(s['end'])}\n{s['text']}\n"
        )
    return "\n".join(blocks)


def _load_job_segments(job_id: str):
    """从内存或历史文件取 segments"""
    job = JOBS.get(job_id)
    if job and job.get("segments"):
        return job
    p = os.path.join(HISTORY_DIR, f"{job_id}.json")
    if os.path.exists(p):
        with open(p, encoding="utf-8") as f:
            return json.load(f)
    return None


@app.get("/api/srt/{job_id}")
def export_srt(job_id: str):
    rec = _load_job_segments(job_id)
    if not rec or not rec.get("segments"):
        return JSONResponse({"error": "任务不存在或无转写结果"}, status_code=404)
    srt = _make_srt(rec["segments"])
    title = (rec.get("title") or "subtitle").strip()[:40]
    safe = "".join(c for c in title if c not in '\\/:*?"<>|') or "subtitle"
    fname = f"{safe}.srt"
    from urllib.parse import quote
    return Response(
        content=srt.encode("utf-8"),
        media_type="application/x-subrip; charset=utf-8",
        headers={"Content-Disposition":
                 f"attachment; filename=\"subtitle.srt\"; "
                 f"filename*=UTF-8''{quote(fname)}"},
    )


# ---------------- 收藏夹批量任务 ----------------
BATCHES = {}         # batch_id -> dict
BATCH_LOCK = threading.Lock()
BATCH_SEM = threading.Semaphore(1)   # 转写是 CPU 密集，串行处理


class BatchRequest(BaseModel):
    text: str             # 收藏夹链接 / 含 fid 的文本
    max_items: int = 30   # 最多处理多少条


def _run_batch(batch_id: str, raw_text: str, max_items: int):
    batch = BATCHES[batch_id]
    workdir = os.path.join(WORK_ROOT, f"batch_{batch_id}")
    os.makedirs(workdir, exist_ok=True)
    try:
        # 1. 解析收藏夹
        fid, _ = extractor.parse_favlist_url(raw_text)
        if not fid:
            raise ValueError("未从内容中解析到收藏夹 ID（fid）。"
                             "请粘贴形如 https://space.bilibili.com/xxx/favlist?fid=123 的链接")
        batch.update(stage="拉取收藏夹列表", progress=5)

        # 2. 拉视频列表
        items = extractor.favlist_items(fid, max_items=max_items)
        if not items:
            raise ValueError("收藏夹是空的，或其中没有视频")
        batch.update(
            total=len(items), stage="开始转写", progress=8,
            items=[{**it, "status": "pending", "job_id": None,
                    "error": None, "segments": None, "full_text": None}
                   for it in items],
        )

        # 3. 逐个处理（串行）
        for i, item in enumerate(batch["items"]):
            if batch.get("canceled"):
                break
            item.update(status="running")
            batch.update(stage=f"转写中 {i+1}/{len(items)}",
                         progress=8 + int(i / len(items) * 90))
            bvid = item["bvid"]
            url = f"https://www.bilibili.com/video/{bvid}"
            try:
                job_id = uuid.uuid4().hex[:12]
                item["job_id"] = job_id
                wav, meta = extractor.bili_extract(url, workdir)
                result = transcriber.transcribe(wav)
                item.update(status="done", title=meta.get("title") or item["title"],
                            duration=meta.get("duration") or item["duration"],
                            segments=result["segments"],
                            full_text=result["full_text"])
                # 存入历史，复用单条的 SRT/txt 导出
                _save_history(job_id, {
                    "job_id": job_id, "url": url, "title": item["title"],
                    "uploader": item.get("upper", ""),
                    "duration": item["duration"],
                    "full_text": result["full_text"],
                    "segments": result["segments"],
                })
            except Exception as e:
                log.warning("批量任务 %s 第 %d 条(%s) 失败: %s", batch_id, i+1, bvid, e)
                item.update(status="error", error=str(e)[:200])
            finally:
                # 清理该条目的临时音频
                for f in os.listdir(workdir):
                    try:
                        os.remove(os.path.join(workdir, f))
                    except OSError:
                        pass
        batch.update(status="done", stage="完成", progress=100,
                     done=sum(1 for x in batch["items"] if x["status"] == "done"))
    except Exception as e:
        log.exception("批量任务 %s 失败", batch_id)
        batch.update(status="error", stage="失败", error=str(e)[:500])
    finally:
        shutil.rmtree(workdir, ignore_errors=True)


@app.post("/api/batch")
def create_batch(req: BatchRequest):
    batch_id = uuid.uuid4().hex[:12]
    with BATCH_LOCK:
        BATCHES[batch_id] = {"id": batch_id, "status": "running", "progress": 0,
                             "stage": "排队中", "created": time.time(),
                             "total": 0, "done": 0, "items": []}
    threading.Thread(target=_run_batch,
                     args=(batch_id, req.text, max(1, min(req.max_items, 100))),
                     daemon=True).start()
    return {"batch_id": batch_id}


@app.get("/api/batch/{batch_id}")
def batch_status(batch_id: str):
    b = BATCHES.get(batch_id)
    if not b:
        return JSONResponse({"error": "批量任务不存在"}, status_code=404)
    # 计算完成数
    b["done"] = sum(1 for x in b["items"] if x["status"] == "done")
    # 返回时去掉超长字段（segments 单独拿）
    slim_items = [{k: v for k, v in it.items()
                   if k not in ("segments", "full_text")}
                  for it in b["items"]]
    return {**{k: v for k, v in b.items() if k != "items"},
            "items": slim_items}


@app.get("/api/batch/{batch_id}/item/{idx}")
def batch_item(batch_id: str, idx: int):
    b = BATCHES.get(batch_id)
    if not b or not (0 <= idx < len(b["items"])):
        return JSONResponse({"error": "不存在"}, status_code=404)
    it = b["items"][idx]
    return {k: v for k, v in it.items()}


# ---------------- Word 文案库批量任务 ----------------
EXPORTS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "exports")
os.makedirs(EXPORTS_DIR, exist_ok=True)


def _fmt_mmss(t: float) -> str:
    t = int(t)
    return f"{t // 60:02d}:{t % 60:02d}"


def _generate_docx(batch: dict) -> str:
    """把批量结果生成 Word 文案库，返回文件路径"""
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    items = batch["items"]
    ok = sum(1 for x in items if x["status"] == "done")

    h = doc.add_heading("视频文案库", level=0)
    p = doc.add_paragraph()
    r = p.add_run(f"生成时间：{time.strftime('%Y-%m-%d %H:%M')}    "
                  f"共 {len(items)} 条 · 转写成功 {ok} 条")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    for i, it in enumerate(items, 1):
        if it["status"] == "done":
            doc.add_heading(f"{i}. {it.get('title') or '(无标题)'}", level=1)
            meta = doc.add_paragraph()
            mr = meta.add_run(
                f"作者：{it.get('uploader') or '-'}   "
                f"时长：{_fmt_mmss(it.get('duration') or 0)}   "
                f"来源：{it.get('url') or ''}")
            mr.font.size = Pt(9)
            mr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
            doc.add_paragraph()  # 空行
            segs = it.get("segments") or []
            if segs:
                for seg in segs:
                    para = doc.add_paragraph()
                    tr = para.add_run(f"[{_fmt_mmss(seg['start'])}] ")
                    tr.bold = True
                    tr.font.size = Pt(10)
                    tr.font.color.rgb = RGBColor(0x4F, 0x7C, 0xFF)
                    para.add_run(seg["text"])
            else:
                wr = doc.add_paragraph()
                wr2 = wr.add_run("（该视频未识别到语音内容，可能是纯音乐/无对白）")
                wr2.font.size = Pt(10)
                wr2.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        else:
            doc.add_heading(f"{i}. {it.get('title') or '(无标题)'}", level=1)
            er = doc.add_paragraph()
            er2 = er.add_run(f"⚠ 转写失败：{it.get('error') or '未知原因'}"
                             f"（链接：{it.get('url') or ''}）")
            er2.font.size = Pt(10)
            er2.font.color.rgb = RGBColor(0xCC, 0x44, 0x44)
        doc.add_paragraph("─" * 40)

    path = os.path.join(EXPORTS_DIR, f"文案库_{batch['id']}.docx")
    doc.save(path)
    return path


def _run_word_batch(batch_id: str, urls: list):
    """Word 上传后的批量转写：全平台链接，逐条 下载→转写→存历史"""
    batch = BATCHES[batch_id]
    workdir = os.path.join(WORK_ROOT, f"batch_{batch_id}")
    os.makedirs(workdir, exist_ok=True)
    try:
        batch.update(total=len(urls), stage="开始转写", progress=5,
                     items=[{"url": u, "title": u.split("//")[-1][:40],
                             "bvid": "", "duration": 0, "upper": "",
                             "status": "pending", "job_id": None,
                             "error": None, "segments": None, "full_text": None}
                            for u in urls])

        for i, item in enumerate(batch["items"]):
            if batch.get("canceled"):
                break
            item.update(status="running")
            batch.update(stage=f"转写中 {i + 1}/{len(urls)}",
                         progress=5 + int(i / len(urls) * 92))
            url = item["url"]
            job_id = uuid.uuid4().hex[:12]
            item["job_id"] = job_id
            try:
                # 与单链接任务完全相同的已验证路径
                cookie_file = get_cookie_file_for(url)
                info = extractor.probe_info(url, cookie_file)
                wav, meta = extractor.download_audio(
                    url, workdir, cookie_file, progress_cb=None)
                result = transcriber.transcribe(wav)
                item.update(status="done",
                            title=meta.get("title") or info.get("title") or url,
                            uploader=meta.get("uploader") or info.get("uploader") or "",
                            duration=meta.get("duration") or info.get("duration") or 0,
                            segments=result["segments"],
                            full_text=result["full_text"])
                _save_history(job_id, {
                    "job_id": job_id, "url": url, "title": item["title"],
                    "uploader": item["uploader"], "duration": item["duration"],
                    "full_text": result["full_text"],
                    "segments": result["segments"],
                })
            except Exception as e:
                log.warning("Word 批量 %s 第 %d 条(%s) 失败: %s",
                            batch_id, i + 1, url, e)
                item.update(status="error", error=str(e)[:200])
            finally:
                for f in os.listdir(workdir):
                    try:
                        os.remove(os.path.join(workdir, f))
                    except OSError:
                        pass

        batch.update(status="done", stage="生成 Word 文案库", progress=99,
                     done=sum(1 for x in batch["items"]
                              if x["status"] == "done"))
        try:
            path = _generate_docx(batch)
            batch.update(stage="完成", progress=100, docx_ready=True,
                         docx_name=os.path.basename(path))
        except Exception as e:
            log.exception("生成 docx 失败 %s", batch_id)
            batch.update(stage="完成（Word 生成失败）", progress=100,
                         docx_ready=False, error=f"Word 生成失败: {e}")
    except Exception as e:
        log.exception("Word 批量任务 %s 失败", batch_id)
        batch.update(status="error", stage="失败", error=str(e)[:500])
    finally:
        shutil.rmtree(workdir, ignore_errors=True)


@app.post("/api/wordbatch")
async def create_word_batch(file: UploadFile = File(...)):
    if not file.filename or not file.filename.lower().endswith(".docx"):
        return JSONResponse({"error": "请上传 .docx 格式的 Word 文档"},
                            status_code=400)
    tmp = os.path.join(WORK_ROOT, f"upload_{uuid.uuid4().hex[:12]}.docx")
    os.makedirs(WORK_ROOT, exist_ok=True)
    try:
        with open(tmp, "wb") as f:
            f.write(await file.read())
        urls = extractor.docx_extract_urls(tmp)
    except Exception as e:
        return JSONResponse({"error": f"Word 文件读取失败：{e}"},
                            status_code=400)
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass

    if not urls:
        return JSONResponse(
            {"error": "文档里没有找到任何链接。支持正文粘贴的链接文本和超链接按钮两种形式。"},
            status_code=400)
    if len(urls) > 100:
        return JSONResponse({"error": f"文档里链接太多（{len(urls)} 条），单次最多 100 条"},
                            status_code=400)

    batch_id = uuid.uuid4().hex[:12]
    with BATCH_LOCK:
        BATCHES[batch_id] = {"id": batch_id, "status": "running", "progress": 0,
                             "stage": "排队中", "created": time.time(),
                             "total": len(urls), "done": 0, "type": "word",
                             "docx_ready": False, "items": []}
    threading.Thread(target=_run_word_batch, args=(batch_id, urls),
                     daemon=True).start()
    return {"batch_id": batch_id, "urls": urls, "count": len(urls)}


@app.get("/api/wordbatch/{batch_id}/docx")
def wordbatch_docx(batch_id: str):
    b = BATCHES.get(batch_id)
    path = os.path.join(EXPORTS_DIR, f"文案库_{batch_id}.docx")
    if not b or not os.path.exists(path):
        return JSONResponse({"error": "文案库不存在或尚未生成"}, status_code=404)
    return FileResponse(
        path,
        media_type=("application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"),
        filename=time.strftime("视频文案库_%Y%m%d_%H%M.docx"),
    )


@app.get("/api/history")
def history_list():
    items = []
    if os.path.isdir(HISTORY_DIR):
        for fn in sorted(os.listdir(HISTORY_DIR), reverse=True):
            if not fn.endswith(".json"):
                continue
            try:
                with open(os.path.join(HISTORY_DIR, fn), encoding="utf-8") as f:
                    r = json.load(f)
                items.append({"job_id": r["job_id"], "time": r["time"],
                              "title": r.get("title", "")[:60],
                              "duration": r.get("duration", 0)})
            except Exception:
                pass
    return {"items": items[:100]}


@app.delete("/api/history/{job_id}")
def history_delete(job_id: str):
    p = os.path.join(HISTORY_DIR, f"{job_id}.json")
    if os.path.exists(p):
        os.remove(p)
    return {"ok": True}


# 静态前端
static_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static")
app.mount("/", StaticFiles(directory=static_dir, html=True), name="static")


if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("PORT", "8300"))
    log.info("启动服务: http://0.0.0.0:%d  （手机连同一 WiFi 访问本机 IP:%d）", port, port)
    uvicorn.run(app, host="0.0.0.0", port=port, log_level="warning")
